"""Convert the blind human-review CSV into a readable Excel workbook."""
import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


def build(csv_path: str, xlsx_path: str) -> None:
    with open(csv_path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    if not rows:
        raise ValueError("The review CSV is empty")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Blind Review"
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "sample_id": 12,
        "model": 24,
        "test_id": 16,
        "attempt": 10,
        "dimension": 24,
        "dimension_description": 48,
        "prompt": 58,
        "response": 85,
        "human_score": 14,
        "human_notes": 34,
    }
    for index, header in enumerate(headers, start=1):
        sheet.column_dimensions[chr(64 + index)].width = widths.get(header, 20)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {
                headers.index("dimension_description") + 1,
                headers.index("prompt") + 1,
                headers.index("response") + 1,
                headers.index("human_notes") + 1,
            })
        row[0].alignment = Alignment(horizontal="center", vertical="top")
        row[8].fill = PatternFill("solid", fgColor="FFF2CC")
        row[9].fill = PatternFill("solid", fgColor="FFF2CC")
        row[0].parent.row_dimensions[row[0].row].height = 55

    score_column = headers.index("human_score") + 1
    score_letter = chr(64 + score_column)
    validation = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    validation.error = "Enter an integer from 1 to 5."
    validation.errorTitle = "Invalid score"
    validation.prompt = "Enter 1, 2, 3, 4, or 5."
    validation.promptTitle = "Human score"
    sheet.add_data_validation(validation)
    validation.add(f"{score_letter}2:{score_letter}{len(rows) + 1}")
    sheet.conditional_formatting.add(
        f"{score_letter}2:{score_letter}{len(rows) + 1}",
        CellIsRule(operator="between", formula=["1", "5"], fill=PatternFill("solid", fgColor="E2F0D9")),
    )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = True

    output = Path(xlsx_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"Wrote {len(rows)} review rows to {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--csv", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    build(args.csv, args.out)
